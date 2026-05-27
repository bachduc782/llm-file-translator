"""スタンドアロン翻訳エンジン — ローカルファイル専用 (.xlsx / .docx / .txt)。"""

from __future__ import annotations

import json
import os
import re
import shutil
import threading
import time
from collections import deque
from concurrent.futures import ThreadPoolExecutor, as_completed
from typing import Callable

import app_config

TRANSLATE_CHUNK = 20    # LLM呼び出し1回あたりの最大アイテム数
CHUNK_CHAR_MAX  = 2000  # 1チャンクあたりの最大文字数
LOCAL_WORKERS   = 40    # ローカルファイル翻訳ワーカー数
FILE_WORKERS    = 4     # 複数ファイルの並列処理数

Progress = Callable[[str], None]

SUPPORTED_EXT = {".xlsx", ".docx", ".txt"}

_LANG_SUFFIX: dict[str, str] = {
    "japanese": "ja", "japanese (日本語)": "ja",
    "english":  "en", "english (英語)":    "en",
    "vietnamese": "vi", "vietnamese (ベトナム語)": "vi",
    "chinese (simplified)": "zh", "chinese (traditional)": "zh-tw",
    "korean": "ko", "french": "fr", "german": "de",
    "spanish": "es", "portuguese": "pt", "thai": "th",
    "indonesian": "id", "malay": "ms", "arabic": "ar",
    "russian": "ru", "italian": "it",
}

_RE_SKIP = re.compile(
    r'^[\d\s,.\-+%/()\[\]{}:;*|]+$'
    r'|^https?://'
    r'|^[a-zA-Z0-9._%+\-]+@\S+'
    r'|^\d{4}[-/]\d{2}[-/]\d{2}',
)


def _is_translatable(text: str) -> bool:
    t = text.strip()
    return len(t) > 2 and not t.startswith("=") and not _RE_SKIP.match(t)


def _noop(*_): pass


def _chunks(items: list, size: int):
    for i in range(0, len(items), size):
        yield items[i : i + size]


def _smart_chunks(items: list) -> list[list]:
    """アイテム数と文字数の両方を考慮してチャンクを分割する。"""
    chunks, cur, cur_chars = [], [], 0
    for item in items:
        text = item[-1] if isinstance(item, tuple) else item
        if cur and (len(cur) >= TRANSLATE_CHUNK or cur_chars + len(text) > CHUNK_CHAR_MAX):
            chunks.append(cur)
            cur, cur_chars = [], 0
        cur.append(item)
        cur_chars += len(text)
    if cur:
        chunks.append(cur)
    return chunks


# ── レートリミッター ──────────────────────────────────────────────────────────

class _RateLimiter:
    def __init__(self, max_calls: int, period: float = 60.0):
        self.max_calls = max_calls
        self.period    = period
        self._lock     = threading.Lock()
        self._calls: deque[float] = deque()

    def acquire(self):
        while True:
            with self._lock:
                now = time.monotonic()
                while self._calls and self._calls[0] <= now - self.period:
                    self._calls.popleft()
                if len(self._calls) < self.max_calls:
                    self._calls.append(now)
                    return
                wait = self._calls[0] + self.period - now + 0.05
            time.sleep(wait)


_limiter: _RateLimiter | None = None
_limiter_key: tuple | None    = None
_limiter_meta_lock             = threading.Lock()


def _get_limiter() -> _RateLimiter:
    global _limiter, _limiter_key
    cfg = app_config.get_llm_config()
    key = (cfg["base_url"], cfg["rate_limit"])
    with _limiter_meta_lock:
        if _limiter is None or _limiter_key != key:
            _limiter     = _RateLimiter(cfg["rate_limit"])
            _limiter_key = key
    return _limiter


# ── LLM ───────────────────────────────────────────────────────────────────────

def _call_llm(prompt: str, retries: int = 3, on_retry=None) -> str:
    from openai import OpenAI

    cfg    = app_config.get_llm_config()
    client = OpenAI(
        base_url=cfg["base_url"],
        api_key=cfg["api_key"],
        default_headers=cfg.get("headers", {}),
        timeout=90.0,
    )
    limiter       = _get_limiter()
    attempt       = 0
    non_429_fails = 0
    last_exc: Exception | None = None

    while True:
        limiter.acquire()
        try:
            resp = client.chat.completions.create(
                model=cfg["model"],
                messages=[{"role": "user", "content": prompt}],
                max_tokens=8192,
            )
            raw = (resp.choices[0].message.content or "").strip()
            if raw.startswith("```"):
                raw = re.sub(r"^```[a-z]*\n?", "", raw)
                raw = re.sub(r"\n?```$", "", raw.strip())
            return raw
        except Exception as exc:
            last_exc = exc
            is_429   = "429" in str(exc) or "rate" in str(exc).lower()
            if is_429:
                wait = min(30 * (attempt + 1), 120)
                if on_retry:
                    on_retry(attempt + 1, "∞", exc, wait)
                time.sleep(wait)
            else:
                non_429_fails += 1
                if non_429_fails >= retries:
                    raise last_exc  # type: ignore[misc]
                wait = 2 ** non_429_fails
                if on_retry:
                    on_retry(non_429_fails, retries, exc, wait)
                time.sleep(wait)
            attempt += 1


def _translate_list(items: list[str], target_language: str, on_retry=None) -> list[str]:
    prompt = (
        f"Translate the following list of text items to {target_language}.\n"
        "Rules:\n"
        "- Translate ALL natural language text, including single words and short phrases.\n"
        "- KEEP AS-IS only items that are NOT natural language: URLs, email addresses, numbers,"
        " and technical tokens (identifiers containing underscores, camelCase, version strings like v1.0,"
        " or mixed letter-digit patterns like ABC123).\n"
        "- Preserve EXACT list length — one output per input.\n"
        "- Return ONLY valid JSON array of strings, no explanation, no markdown.\n\n"
        f"Input:\n{json.dumps(items, ensure_ascii=False)}"
    )
    raw = _call_llm(prompt, on_retry=on_retry)
    if not raw.strip():
        return items
    try:
        result = json.loads(raw)
    except json.JSONDecodeError:
        from json_repair import repair_json
        repaired = repair_json(raw)
        if not repaired or repaired in ("null", "[]", "{}"):
            return items
        try:
            result = json.loads(repaired)
        except json.JSONDecodeError:
            return items
    if not isinstance(result, list):
        return items
    return (result + items[len(result):])[:len(items)]


# ── 重複排除翻訳 ──────────────────────────────────────────────────────────────

def _translate_unique(
    texts: list[str],
    target_language: str,
    progress: Progress,
) -> dict[str, str]:
    """重複を排除してから翻訳し、{原文: 訳文} のキャッシュを返す。"""
    unique = list(dict.fromkeys(texts))
    dupes  = len(texts) - len(unique)
    if dupes:
        progress(f"{len(unique)}件のユニークテキストを翻訳 ({dupes}件重複スキップ)")

    cache: dict[str, str] = {}
    lock         = threading.Lock()
    done         = 0
    all_chunks   = _smart_chunks([(i, t) for i, t in enumerate(unique)])
    total_chunks = len(all_chunks)

    def _worker(chunk: list, idx: int):
        nonlocal done

        def _on_retry(attempt, total_r, exc, wait):
            progress(f"リトライ [{idx}/{total_chunks}] 試行 {attempt}/{total_r} (待機 {wait}秒) — {exc}")

        translated = _translate_list([t for _, t in chunk], target_language, on_retry=_on_retry)
        with lock:
            for (_, orig), new in zip(chunk, translated):
                cache[orig] = new
            done += len(chunk)
            progress(f"[{idx}/{total_chunks}] 完了 — {done}/{len(unique)}件翻訳済み")

    with ThreadPoolExecutor(max_workers=LOCAL_WORKERS) as pool:
        futures = {pool.submit(_worker, chunk, i + 1): i for i, chunk in enumerate(all_chunks)}
        for fut in as_completed(futures):
            try:
                fut.result()
            except Exception as e:
                progress(f"ERROR チャンク {futures[fut] + 1}: {e}")

    return cache


# ── ファイル名生成 ─────────────────────────────────────────────────────────────

def _clone_path(file_path: str, target_language: str) -> str:
    base, ext = os.path.splitext(file_path)
    suffix = _LANG_SUFFIX.get(target_language.lower(), target_language[:2].lower())
    return f"{base}-{suffix}{ext}"


# ── Excel (.xlsx) ──────────────────────────────────────────────────────────────

def _translate_xlsx(file_path: str, sheet_name: str, target_language: str, progress: Progress) -> int:
    import openpyxl

    wb = openpyxl.load_workbook(file_path)

    if sheet_name:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"シート '{sheet_name}' が見つかりません。利用可能: {wb.sheetnames}")
        ws_list = [wb[sheet_name]]
    else:
        ws_list = [wb[sn] for sn in wb.sheetnames]

    all_cells: list[tuple[str, int, int, str]] = []
    total_raw = 0
    for ws in ws_list:
        progress(f"スキャン中 [{ws.title}]…")
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.strip():
                    total_raw += 1
                    if _is_translatable(cell.value):
                        all_cells.append((ws.title, cell.row, cell.column, cell.value))

    skipped = total_raw - len(all_cells)
    progress(
        f"{len(ws_list)}シートで{len(all_cells)}セルを翻訳します"
        + (f" ({skipped}件スキップ)" if skipped else "")
    )
    if not all_cells:
        return 0

    cache = _translate_unique([t for _, _, _, t in all_cells], target_language, progress)
    translated_map = {
        (sn, r, c): cache[orig]
        for sn, r, c, orig in all_cells
        if cache.get(orig, orig) != orig
    }

    for (sn, r, c), new in translated_map.items():
        wb[sn].cell(row=r, column=c).value = new

    progress("ファイルを保存中…")
    wb.save(file_path)
    return len(all_cells)


# ── Word (.docx) ───────────────────────────────────────────────────────────────

def _set_para_text(para, new_text: str):
    if not para.runs:
        para.add_run(new_text)
        return
    para.runs[0].text = new_text
    for run in para.runs[1:]:
        run.text = ""


def _iter_doc_paragraphs(doc):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def _translate_docx(file_path: str, target_language: str, progress: Progress) -> int:
    from docx import Document

    doc = Document(file_path)

    para_items: list[tuple[object, str]] = [
        (para, para.text.strip())
        for para in _iter_doc_paragraphs(doc)
        if para.text.strip() and _is_translatable(para.text.strip())
    ]
    total = len(para_items)
    progress(f"{total}段落を翻訳します")
    if not total:
        return 0

    cache = _translate_unique([text for _, text in para_items], target_language, progress)

    for i, (para, orig) in enumerate(para_items):
        new = cache.get(orig, orig)
        if orig != new:
            _set_para_text(para, new)

    progress("ファイルを保存中…")
    doc.save(file_path)
    return total


# ── プレーンテキスト (.txt) ────────────────────────────────────────────────────

def _translate_txt(file_path: str, target_language: str, progress: Progress) -> int:
    with open(file_path, encoding="utf-8", errors="replace") as f:
        lines = f.readlines()

    indexed = [
        (i, line.rstrip("\r\n"))
        for i, line in enumerate(lines)
        if line.strip() and _is_translatable(line.strip())
    ]
    total = len(indexed)
    progress(f"{total}行を翻訳します")
    if not total:
        return 0

    cache = _translate_unique([t for _, t in indexed], target_language, progress)

    result_lines = list(lines)
    for i, orig_text in indexed:
        new_text = cache.get(orig_text, orig_text)
        orig     = lines[i]
        ending   = "\r\n" if orig.endswith("\r\n") else "\n"
        result_lines[i] = new_text + ending

    progress("ファイルを保存中…")
    with open(file_path, "w", encoding="utf-8") as f:
        f.writelines(result_lines)
    return total


# ── 公開API ───────────────────────────────────────────────────────────────────

def translate_local_file(
    file_path: str,
    target_language: str,
    sheet_name: str = "",
    progress: Progress = _noop,
) -> dict:
    """ローカルファイルを1つ複製して翻訳する。結果dictを返す。"""
    if not os.path.isfile(file_path):
        return {"error": f"ファイルが見つかりません: {file_path}"}

    ext = os.path.splitext(file_path)[1].lower()
    if ext not in SUPPORTED_EXT:
        return {"error": f"未対応の形式: {ext}。対応形式: {', '.join(sorted(SUPPORTED_EXT))}"}

    clone         = _clone_path(file_path, target_language)
    clone_name    = os.path.basename(clone)
    original_name = os.path.basename(file_path)

    progress("ファイルを複製中…")
    try:
        shutil.copy2(file_path, clone)
    except Exception as e:
        return {"error": f"ファイルを複製できません: {e}"}
    progress(f"複製完了: {original_name} → {clone_name}")

    try:
        if ext == ".xlsx":
            total = _translate_xlsx(clone, sheet_name, target_language, progress)
            return {"ok": True, "clone_name": clone_name, "clone_path": clone,
                    "file_type": "xlsx", "cells_translated": total}

        elif ext == ".docx":
            total = _translate_docx(clone, target_language, progress)
            return {"ok": True, "clone_name": clone_name, "clone_path": clone,
                    "file_type": "docx", "paragraphs_translated": total}

        else:
            total = _translate_txt(clone, target_language, progress)
            return {"ok": True, "clone_name": clone_name, "clone_path": clone,
                    "file_type": "txt", "lines_translated": total}

    except Exception as e:
        return {"error": str(e), "clone_name": clone_name, "clone_path": clone}


def translate_local_files(
    file_paths: list[str],
    target_language: str,
    sheet_name: str = "",
    progress: Progress = _noop,
) -> dict:
    """複数のローカルファイルを並列で複製・翻訳する。"""
    results: list[dict] = [{}] * len(file_paths)

    def _process(idx: int, fp: str):
        def _prog(msg):
            progress(f"[{idx + 1}/{len(file_paths)}] {msg}")
        results[idx] = translate_local_file(fp, target_language, sheet_name, _prog)

    with ThreadPoolExecutor(max_workers=min(len(file_paths), FILE_WORKERS)) as pool:
        futures = {pool.submit(_process, i, fp): i for i, fp in enumerate(file_paths)}
        for fut in as_completed(futures):
            fut.result()

    succeeded = sum(1 for r in results if r.get("ok"))
    return {
        "total": len(file_paths),
        "succeeded": succeeded,
        "failed": len(file_paths) - succeeded,
        "results": results,
    }
