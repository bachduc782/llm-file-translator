"""ローカルファイル翻訳エンジン — Excel (.xlsx)、Word (.docx)、テキスト (.txt)。"""

from __future__ import annotations

import io
import os
import re
import shutil
import threading
import zipfile
from concurrent.futures import ThreadPoolExecutor, as_completed

from src.tools.translate_google import (
    _translate_list, _is_translatable, _force_translate, _chunks,
    TRANSLATE_CHUNK, FILE_WORKERS,
    Progress, _noop,
)

# ローカルファイルはGoogle APIレート制限なし — ワーカー数=レート上限に設定
# レートリミッタースロットが空いた瞬間に常にワーカーが待機している状態にする
LOCAL_WORKERS = 40
from src.tools.google_drive import _LANG_SUFFIX

SUPPORTED_EXT = {".xlsx", ".docx", ".txt"}

def _clone_path(file_path: str, target_language: str) -> str:
    base, ext = os.path.splitext(file_path)
    suffix = _LANG_SUFFIX.get(target_language.lower(), target_language[:2].lower())
    return f"{base}-{suffix}{ext}"


# ── Excel (.xlsx) ──────────────────────────────────────────────────────────────

from xml.sax.saxutils import escape as _xml_escape

# sharedStrings.xml の解析・書き換え用
_SI_RE         = re.compile(r'<si>.*?</si>', re.S)
_RPH_RE        = re.compile(r'<rPh\b[^>]*>.*?</rPh>', re.S)   # 振り仮名（ルビ）
_PHON_RE       = re.compile(r'<phoneticPr\b[^>]*/>')
_T_INNER_RE    = re.compile(r'<t\b[^>]*>(.*?)</t>', re.S)
_NUM_REF_RE    = re.compile(r'&#(\d+);')
_HEX_REF_RE    = re.compile(r'&#x([0-9A-Fa-f]+);')
# 共有文字列を参照するセル: <c ... t="s" ...><v>IDX</v>
_SHARED_REF_RE = re.compile(r'<c\b[^>]*\bt="s"[^>]*>\s*<v>(\d+)</v>')


def _xml_unescape(s: str) -> str:
    s = s.replace('&lt;', '<').replace('&gt;', '>')
    s = s.replace('&quot;', '"').replace('&apos;', "'")
    s = _NUM_REF_RE.sub(lambda m: chr(int(m.group(1))), s)
    s = _HEX_REF_RE.sub(lambda m: chr(int(m.group(1), 16)), s)
    return s.replace('&amp;', '&')  # &amp; は最後にデコード


def _si_plain_text(si: str) -> str:
    """<si> ブロックのプレーンテキストを取得する（振り仮名は除外）。"""
    body = _PHON_RE.sub('', _RPH_RE.sub('', si))
    return _xml_unescape(''.join(_T_INNER_RE.findall(body)))


def _xlsx_sheet_paths(zip_data: dict) -> list[tuple[str, str]]:
    """workbook.xml から (シート名, シートのzip内パス) のリストを返す。"""
    wb_xml = zip_data.get("xl/workbook.xml", b"").decode("utf-8", "replace")
    rels   = zip_data.get("xl/_rels/workbook.xml.rels", b"").decode("utf-8", "replace")
    rid_to_target: dict[str, str] = {}
    for m in re.finditer(r'<Relationship\b[^>]*/>', rels):
        el = m.group(0)
        i = re.search(r'Id="([^"]+)"', el)
        t = re.search(r'Target="([^"]+)"', el)
        if i and t:
            rid_to_target[i.group(1)] = t.group(1)
    result: list[tuple[str, str]] = []
    for m in re.finditer(r'<sheet\b[^>]*/>', wb_xml):
        el  = m.group(0)
        nm  = re.search(r'\bname="([^"]+)"', el)
        rid = re.search(r'r:id="([^"]+)"', el)
        if not (nm and rid):
            continue
        tgt = rid_to_target.get(rid.group(1))
        if not tgt:
            continue
        path = tgt[1:] if tgt.startswith("/") else ("xl/" + tgt)
        result.append((_xml_unescape(nm.group(1)), path))
    return result


def _translate_xlsx(file_path: str, sheet_name: str, target_language: str, progress: Progress) -> int:
    """
    原本 zip をそのままコピーし、xl/sharedStrings.xml 内のテキストのみ翻訳して差し替える。
    drawings・charts・vml・comments・styles・worksheets はバイト単位で保持され、
    数式も破壊されない（openpyxl 保存による図形消失・破損を完全に回避する）。
    """
    from collections import Counter

    with zipfile.ZipFile(file_path, "r") as zf:
        zip_data = {n: zf.read(n) for n in zf.namelist()}

    SS = "xl/sharedStrings.xml"
    if SS not in zip_data:
        progress("翻訳対象の共有文字列がありません")
        return 0
    ss_xml = zip_data[SS].decode("utf-8")

    # フェーズ1：sharedStrings の <si> を抽出
    si_texts = [_si_plain_text(m.group(0)) for m in _SI_RE.finditer(ss_xml)]

    # 翻訳スコープ（シート指定があればそのシートのみ）
    sheets = _xlsx_sheet_paths(zip_data)
    if sheet_name:
        scope = [(nm, p) for nm, p in sheets if nm == sheet_name]
        if not scope:
            raise ValueError(
                f"シート '{sheet_name}' が見つかりません。利用可能: {[n for n, _ in sheets]}"
            )
    else:
        scope = sheets

    # スコープ内シートを走査し、参照されている共有文字列インデックスを集計
    used: Counter = Counter()
    for nm, path in scope:
        progress(f"スキャン中 [{nm}]…")
        data = zip_data.get(path, b"").decode("utf-8", "replace")
        for m in _SHARED_REF_RE.finditer(data):
            used[int(m.group(1))] += 1

    # 翻訳候補（翻訳可能な文字列のみ）
    candidates: list[tuple[int, str]] = []   # (si_index, stripped_text)
    skipped_cells = 0
    for idx, cnt in used.items():
        if idx >= len(si_texts):
            continue
        text = si_texts[idx].strip()
        if not text:
            continue
        if _is_translatable(text):
            candidates.append((idx, text))
        else:
            skipped_cells += cnt

    if not candidates:
        progress("翻訳対象の文字列がありません")
        return 0

    # フェーズ2：dedup
    texts  = [t for _, t in candidates]
    counts = Counter(texts)
    unique = list(dict.fromkeys(texts))
    occ_total = sum(used[idx] for idx, _ in candidates)
    progress(
        f"{len(scope)}シートで{len(unique)}件のユニークテキストを翻訳します"
        + (f" ({skipped_cells}セルスキップ)" if skipped_cells else "")
    )
    for text, cnt in sorted(counts.items(), key=lambda x: -x[1]):
        if cnt > 1:
            preview = text[:60].replace("\n", "↵")
            progress(f"  重複 x{cnt}: {preview}{'…' if len(text) > 60 else ''}")

    # フェーズ3：翻訳（uniqueのみ・並列）
    cache        : dict[str, str] = {}
    lock         = threading.Lock()
    done         = 0
    all_chunks   = list(_chunks(list(enumerate(unique)), TRANSLATE_CHUNK))
    total_chunks = len(all_chunks)

    def _worker(chunk: list, idx: int):
        nonlocal done
        progress(f"[{idx}/{total_chunks}] LLMを呼び出し中 ({len(chunk)}件)…")

        def _on_retry(attempt, total_r, exc, wait):
            progress(f"リトライ [{idx}/{total_chunks}] 試行 {attempt}/{total_r} (待機 {wait}秒) — {exc}")

        texts_chunk = [t for _, t in chunk]
        translated  = _translate_list(texts_chunk, target_language, on_retry=_on_retry)
        with lock:
            for (_, orig), new in zip(chunk, translated):
                cache[orig] = new
            done += len(chunk)
            progress(f"[{idx}/{total_chunks}] 完了 — {done}/{len(unique)}件翻訳済み")

    with ThreadPoolExecutor(max_workers=LOCAL_WORKERS) as pool:
        futures = {pool.submit(_worker, chunk, i + 1): i for i, chunk in enumerate(all_chunks)}
        for fut in as_completed(futures):
            try:
                fut.result()
            except Exception as e:
                progress(f"ERROR チャンク {futures[fut] + 1}: {e}")

    # フェーズ4a：LLMが同一テキストを返したユニーク分を強制再翻訳（並列）
    retry_texts = list({t for t in unique if cache.get(t, t) == t})
    if retry_texts:
        progress(f"同一テキスト返却 {len(retry_texts)}件を強制再翻訳…")
        force_done = 0

        def _force_worker(text: str):
            nonlocal force_done
            forced = _force_translate(text, target_language)
            with lock:
                if forced != text:
                    cache[text] = forced
                force_done += 1
                progress(f"強制再翻訳 {force_done}/{len(retry_texts)}件完了")

        with ThreadPoolExecutor(max_workers=LOCAL_WORKERS) as pool:
            futures = [pool.submit(_force_worker, t) for t in retry_texts]
            for fut in as_completed(futures):
                try:
                    fut.result()
                except Exception as e:
                    progress(f"ERROR 強制再翻訳: {e}")

    # フェーズ4b：sharedStrings を書き換え（翻訳されたインデックスのみ）
    idx_to_new = {
        idx: cache[t] for idx, t in candidates if cache.get(t, t) != t
    }
    if not idx_to_new:
        progress("⚠ 翻訳結果が原文と同一のため、更新する文字列はありません")
        return 0

    _counter = {"i": -1}

    def _repl(m):
        _counter["i"] += 1
        new = idx_to_new.get(_counter["i"])
        if new is None:
            return m.group(0)
        return f'<si><t xml:space="preserve">{_xml_escape(new)}</t></si>'

    zip_data[SS] = _SI_RE.sub(_repl, ss_xml).encode("utf-8")

    changed_occ = sum(used[idx] for idx in idx_to_new)
    unchanged   = occ_total - changed_occ
    if unchanged:
        progress(f"⚠ {unchanged}セルはLLMが同一テキストを返したため未更新")

    # フェーズ5：原本構造のまま再パッケージ（sharedStrings 以外はバイト保持）
    progress("ファイルを保存中…")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as out_zip:
        for name, data in zip_data.items():
            out_zip.writestr(name, data)
    with open(file_path, "wb") as f:
        f.write(buf.getvalue())

    return changed_occ


# ── Word (.docx) ───────────────────────────────────────────────────────────────

def _set_para_text(para, new_text: str):
    """段落テキストを置換する。先頭テキストランのフォーマットを維持し、画像ランを保護する。"""
    from docx.oxml.ns import qn

    # w:drawing / w:object / w:pict を含むランは画像ランなので変更しない
    _IMAGE_TAGS = {qn("w:drawing"), qn("w:object"), qn("w:pict")}

    def _has_image(run):
        return any(child.tag in _IMAGE_TAGS for child in run._r)

    text_runs = [r for r in para.runs if not _has_image(r)]

    if not text_runs:
        para.add_run(new_text)
        return
    text_runs[0].text = new_text
    for run in text_runs[1:]:
        run.text = ""


def _iter_doc_paragraphs(doc):
    """本文とテーブルから全段落を生成する。"""
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def _translate_docx(file_path: str, target_language: str, progress: Progress) -> int:
    from docx import Document

    doc = Document(file_path)

    para_items: list[tuple[object, str]] = [
        (para, para.text.strip())
        for para in _iter_doc_paragraphs(doc)
        if para.text.strip() and _is_translatable(para.text.strip())
    ]
    total = len(para_items)
    progress(f"{total}段落を翻訳します")
    if not total:
        return 0

    translated_results: dict[int, str] = {}
    lock         = threading.Lock()
    done         = 0
    indexed      = [(i, text) for i, (_, text) in enumerate(para_items)]
    all_chunks   = list(_chunks(indexed, TRANSLATE_CHUNK))
    total_chunks = len(all_chunks)

    def _worker(chunk: list, idx: int):
        nonlocal done
        progress(f"[{idx}/{total_chunks}] LLMを呼び出し中 ({len(chunk)}段落)…")

        def _on_retry(attempt, total_r, exc, wait):
            progress(f"リトライ [{idx}/{total_chunks}] 試行 {attempt}/{total_r} (待機 {wait}秒) — {exc}")

        texts      = [t for _, t in chunk]
        translated = _translate_list(texts, target_language, on_retry=_on_retry)
        with lock:
            for (i, orig), new in zip(chunk, translated):
                if orig != new:
                    translated_results[i] = new
            done += len(chunk)
            progress(f"[{idx}/{total_chunks}] 完了 — {done}/{total}段落翻訳済み")

    with ThreadPoolExecutor(max_workers=LOCAL_WORKERS) as pool:
        futures = {pool.submit(_worker, chunk, i + 1): i for i, chunk in enumerate(all_chunks)}
        for fut in as_completed(futures):
            try:
                fut.result()
            except Exception as e:
                progress(f"ERROR チャンク {futures[fut] + 1}: {e}")

    for i, new_text in translated_results.items():
        _set_para_text(para_items[i][0], new_text)

    progress("ファイルを保存中…")
    doc.save(file_path)
    return total


# ── プレーンテキスト (.txt) ────────────────────────────────────────────────────

def _translate_txt_local(file_path: str, target_language: str, progress: Progress) -> int:
    with open(file_path, encoding="utf-8", errors="replace") as f:
        lines = f.readlines()

    indexed = [
        (i, line.rstrip("\r\n"))
        for i, line in enumerate(lines)
        if line.strip() and _is_translatable(line.strip())
    ]
    total = len(indexed)
    progress(f"{total}行を翻訳します")
    if not total:
        return 0

    translated_map: dict[int, str] = {}
    lock         = threading.Lock()
    done         = 0
    all_chunks   = list(_chunks(indexed, TRANSLATE_CHUNK))
    total_chunks = len(all_chunks)

    def _worker(chunk: list, idx: int):
        nonlocal done
        progress(f"[{idx}/{total_chunks}] LLMを呼び出し中 ({len(chunk)}行)…")

        def _on_retry(attempt, total_r, exc, wait):
            progress(f"リトライ [{idx}/{total_chunks}] 試行 {attempt}/{total_r} (待機 {wait}秒) — {exc}")

        texts      = [t for _, t in chunk]
        translated = _translate_list(texts, target_language, on_retry=_on_retry)
        with lock:
            for (i, _), new in zip(chunk, translated):
                translated_map[i] = new
            done += len(chunk)
            progress(f"[{idx}/{total_chunks}] 完了 — {done}/{total}行翻訳済み")

    with ThreadPoolExecutor(max_workers=LOCAL_WORKERS) as pool:
        futures = {pool.submit(_worker, chunk, i + 1): i for i, chunk in enumerate(all_chunks)}
        for fut in as_completed(futures):
            try:
                fut.result()
            except Exception as e:
                progress(f"ERROR チャンク {futures[fut] + 1}: {e}")

    result_lines = list(lines)
    for i, new_text in translated_map.items():
        orig    = lines[i]
        ending  = "\r\n" if orig.endswith("\r\n") else "\n"
        result_lines[i] = new_text + ending

    progress("ファイルを保存中…")
    with open(file_path, "w", encoding="utf-8") as f:
        f.writelines(result_lines)
    return total


# ── 公開API ───────────────────────────────────────────────────────────────────

def translate_local_file(
    file_path: str,
    target_language: str,
    sheet_name: str = "",
    progress: Progress = _noop,
) -> dict:
    """ローカルファイルを1つ複製して翻訳する。結果dictを返す。"""
    if not os.path.isfile(file_path):
        return {"error": f"ファイルが見つかりません: {file_path}"}

    ext = os.path.splitext(file_path)[1].lower()
    if ext not in SUPPORTED_EXT:
        return {"error": f"未対応の形式: {ext}。対応形式: {', '.join(sorted(SUPPORTED_EXT))}"}

    clone = _clone_path(file_path, target_language)
    clone_name    = os.path.basename(clone)
    original_name = os.path.basename(file_path)

    progress("ファイルを複製中…")
    try:
        shutil.copy2(file_path, clone)
    except Exception as e:
        return {"error": f"ファイルを複製できません: {e}"}
    progress(f"複製完了: {original_name} → {clone_name}")

    try:
        if ext == ".xlsx":
            total = _translate_xlsx(clone, sheet_name, target_language, progress)
            return {"ok": True, "clone_name": clone_name, "clone_path": clone,
                    "file_type": "xlsx", "cells_translated": total}

        elif ext == ".docx":
            total = _translate_docx(clone, target_language, progress)
            return {"ok": True, "clone_name": clone_name, "clone_path": clone,
                    "file_type": "docx", "paragraphs_translated": total}

        else:  # .txt
            total = _translate_txt_local(clone, target_language, progress)
            return {"ok": True, "clone_name": clone_name, "clone_path": clone,
                    "file_type": "txt", "lines_translated": total}

    except Exception as e:
        return {"error": str(e), "clone_name": clone_name, "clone_path": clone}


def translate_local_files(
    file_paths: list[str],
    target_language: str,
    sheet_name: str = "",
    progress: Progress = _noop,
) -> dict:
    """複数のローカルファイルを並列で複製・翻訳する。"""
    results: list[dict] = [{}] * len(file_paths)

    def _process(idx: int, fp: str):
        def _prog(msg):
            progress(f"[{idx + 1}/{len(file_paths)}] {msg}")
        results[idx] = translate_local_file(fp, target_language, sheet_name, _prog)

    with ThreadPoolExecutor(max_workers=min(len(file_paths), FILE_WORKERS)) as pool:
        futures = {pool.submit(_process, i, fp): i for i, fp in enumerate(file_paths)}
        for fut in as_completed(futures):
            fut.result()

    succeeded = sum(1 for r in results if r.get("ok"))
    return {
        "total": len(file_paths),
        "succeeded": succeeded,
        "failed": len(file_paths) - succeeded,
        "results": results,
    }
